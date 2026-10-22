#!/usr/bin/env python3
"""Publish a verified, versioned Boverket climate-material catalog from official XLSX.

The XLSX is an official machine-readable distribution of the same source data
as Boverket's climate database API. Until the exact OpenAPI route is verified,
do not guess a URL or silently substitute malformed data.
"""
import io,json,os,urllib.request
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook

URL="https://www.boverket.se/contentassets/4668ed4cc3da447385788ed30bff7d49/boverkets-klimatdatabas-version-02.07.000-sv-se.xlsx"
COLUMNS={"id":"Resurs-ID","name":"Produktnamn","category":"Kategori","version":"Version",
"unit":"Enhet för klimatpåverkan","typical":"A1-A3 byggproduktens klimatpåverkan GWP-GHG, typiskt värde",
"conservative":"A1-A3 klimatpåverkan, konservativt värde",
"a4":"A4 transporters klimatpåverkan GWP-GHG","a5":"A5 byggspills klimatpåverkan GWP-GHG, konservativt värde"}
def num(x):
    return x if isinstance(x,(int,float)) and not isinstance(x,bool) else None
def main():
    req=urllib.request.Request(URL,headers={"User-Agent":"Lulea-statistik climate materials catalog/1.0"})
    with urllib.request.urlopen(req,timeout=65) as f: content=f.read()
    wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    ws=wb.active
    it=ws.iter_rows(values_only=True)
    head=[str(x).strip() if x is not None else "" for x in next(it)]
    missing=[c for c in COLUMNS.values() if c not in head]
    if missing: raise ValueError("Boverket schema changed: "+repr(missing))
    ix={k:head.index(v) for k,v in COLUMNS.items()}
    data=[]
    for row in it:
        if not row[ix["id"]] or not row[ix["name"]]:continue
        material={k:row[i] for k,i in ix.items()}
        for k in ("typical","conservative","a4","a5"):material[k]=num(material[k])
        material["id"]=str(material["id"])
        data.append(material)
    if len(data)<180: raise ValueError("Too few data rows; refusing to publish")
    versions={m["version"] for m in data}
    if len(versions)!=1: raise ValueError("Mixed climate dataset versions")
    out=Path("docs/data/materials.json");out.parent.mkdir(parents=True,exist_ok=True)
    obj={"source":"Boverket","source_url":URL,"version":next(iter(versions)),
    "retrieved_at_utc":datetime.now(timezone.utc).isoformat(),"count":len(data),"materials":data}
    new=json.dumps(obj,ensure_ascii=False,indent=2,default=str)+"\n"
    # Avoid daily commit churn if no source values changed: keep existing retrieved timestamp.
    if out.exists():
        old=json.loads(out.read_text(encoding="utf8"))
        if old.get("materials")==data and old.get("version")==obj["version"]:
            print("No source changes",len(data));return
    out.write_text(new,encoding="utf8")
    print("Published",len(data),"Boverket resources",obj["version"])
if __name__=="__main__":main()
